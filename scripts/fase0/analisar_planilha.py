#!/usr/bin/env python3
"""Fase 0 - Analise da planilha de controle de NFS-e / P.Os da Ameta.

Le a planilha original (sem altera-la) e gera:
  - um resumo no terminal (abas, colunas, tipos, % de vazios, status, duplicidades);
  - inconsistencias.csv: uma linha por problema encontrado, com o numero da linha na planilha.

Uso:
    python3 scripts/fase0/analisar_planilha.py data/Controle_AMETA_4.xlsx [pasta_saida]

A planilha contem dados de clientes e fica fora do Git (pasta data/ no .gitignore).
Dependencia: openpyxl (pip install openpyxl).
"""
import collections
import csv
import datetime
import os
import re
import sys

import openpyxl

ABA_PRINCIPAL = "AMETA-2025-NEW"

# Mapeamento proposto (ver docs/fase-0-analise.md). None = aguardando decisao do usuario.
MAPA_STATUS = {
    "AGUARDANDO LIBERAÇÃO": ("AGUARDANDO_LIBERACAO", False),
    "EMITIR NOTA": ("EMITIR_NOTA", False),
    "EM EXECUÇÃO": ("EM_EXECUCAO", False),
    "EMITIDA NOTA FISCAL": ("EMITIDA", False),
    "EMITIDA NOTA FISCAL C/ MULTA": ("EMITIDA", True),
    "P.O CANCELADO": ("CANCELADA", False),
    "PEDIDO CANCELADO": ("CANCELADA", False),
    "SITE CANCELADO": ("CANCELADA", False),
    "NOTA DUPLICADA CANCELADA": None,
    "PENDENTE": None,
}


def so_digitos(v):
    return None if v is None else re.sub(r"\D", "", str(v)) or None


def chave_projeto(v):
    """Normaliza variacoes de acentuacao quebrada (S?o, S#o, S�o, SÃo...)."""
    if v is None:
        return None
    s = str(v).strip().lower()
    s = re.sub(r"s.o paulo", "sao paulo", s)
    return re.sub(r"\W", "", s)


def vazio(v):
    return v is None or (isinstance(v, str) and v.strip() in ("", ".", "-", "_"))


def carregar(caminho):
    wf = openpyxl.load_workbook(caminho)  # formulas
    wv = openpyxl.load_workbook(caminho, data_only=True)  # valores calculados
    return wf, wv


def main():
    caminho = sys.argv[1] if len(sys.argv) > 1 else "data/Controle_AMETA_4.xlsx"
    saida = sys.argv[2] if len(sys.argv) > 2 else "data/fase0"
    os.makedirs(saida, exist_ok=True)
    wf, wv = carregar(caminho)

    print("== Abas")
    for ws in wf.worksheets:
        print(f"  {ws.title!r}: {ws.max_row} linhas x {ws.max_column} colunas, tabelas={list(ws.tables)}")

    sf, sv = wf[ABA_PRINCIPAL], wv[ABA_PRINCIPAL]
    cab = [c.value for c in sf[1]]
    linhas = []
    for rf, rv in zip(sf.iter_rows(min_row=2), sv.iter_rows(min_row=2)):
        d = {"_linha": rf[0].row}
        for i, (cf, cv) in enumerate(zip(rf, rv)):
            nome = (cab[i] or f"(sem nome {cf.column_letter})").strip()
            d[nome] = cv.value
            d[nome + "#formula"] = isinstance(cf.value, str) and cf.value.startswith("=")
        linhas.append(d)
    colunas = [(c or f"(sem nome {openpyxl.utils.get_column_letter(i + 1)})").strip() for i, c in enumerate(cab)]
    n = len(linhas)
    print(f"\n== {ABA_PRINCIPAL}: {n} linhas de dados")
    print(f"  {'coluna':22} {'vazios':>7} {'%':>6} {'formulas':>8} {'distintos':>9}  tipos")
    for c in colunas:
        vals = [r[c] for r in linhas]
        nv = sum(vazio(v) for v in vals)
        nf = sum(r[c + "#formula"] for r in linhas)
        tipos = collections.Counter(type(v).__name__ for v in vals if not vazio(v))
        print(f"  {c:22} {nv:7} {100 * nv / n:5.1f}% {nf:8} {len(set(map(str, vals))):9}  {dict(tipos)}")

    st = lambda r: r["STATUS"]
    print("\n== Status")
    for s, q in collections.Counter(map(st, linhas)).most_common():
        print(f"  {q:6}  {s!r:34} -> {MAPA_STATUS.get(s, 'NAO RECONHECIDO')}")

    problemas = []

    def p(r, tipo, detalhe=""):
        problemas.append({"linha": r["_linha"], "tipo": tipo, "po": r["P.O"], "item": r["ITEM"],
                          "status": st(r), "nfse": r["N°NFS-e"], "detalhe": detalhe})

    emitida = lambda r: st(r).startswith("EMITIDA NOTA FISCAL")
    for r in linhas:
        po = r["P.O"]
        if po is None:
            p(r, "PO_VAZIA")
        elif isinstance(po, str) and not re.fullmatch(r"\d+", po):
            p(r, "PO_COM_CARACTERES_EXTRAS", repr(po))
        if po is not None and len(so_digitos(po) or "") != 10:
            p(r, "PO_TAMANHO_DIFERENTE_DE_10", repr(po))
        if MAPA_STATUS.get(st(r), "x") is None:
            p(r, "STATUS_SEM_MAPEAMENTO", st(r))
        if not isinstance(r["ITEM"], int):
            p(r, "ITEM_VAZIO_OU_MULTIPLO", repr(r["ITEM"]))
        nf = r["N°NFS-e"]
        if nf is not None and not re.fullmatch(r"\d+", str(nf)):
            p(r, "NFSE_NAO_NUMERICA", repr(nf))
        if emitida(r) and (nf is None or r["NOTA EMITIDA"] is None):
            p(r, "EMITIDA_SEM_NFSE_OU_DATA", f"nfse={nf} data={r['NOTA EMITIDA']}")
        if st(r) in ("AGUARDANDO LIBERAÇÃO", "PENDENTE", "PEDIDO CANCELADO", "P.O CANCELADO") and nf is not None:
            p(r, "NFSE_PREENCHIDA_EM_STATUS_NAO_EMITIDO", str(nf))
        dt = r["NOTA EMITIDA"]
        if isinstance(dt, datetime.datetime) and dt.year < 2020:
            p(r, "DATA_EMISSAO_SUSPEITA", dt.date().isoformat())
        multa = r["MULTA"]
        if st(r) == "EMITIDA NOTA FISCAL C/ MULTA" and multa is None:
            p(r, "STATUS_COM_MULTA_SEM_PERCENTUAL")
        if st(r) == "EMITIDA NOTA FISCAL" and multa is not None:
            p(r, "PERCENTUAL_MULTA_EM_STATUS_SEM_MULTA", str(multa))
        if r["MULTA#formula"]:
            p(r, "FORMULA_NA_COLUNA_MULTA")
        if r["PREÇO ORIGINAL"] is None:
            p(r, "PRECO_ORIGINAL_VAZIO")
        if vazio(r["FASE"]):
            p(r, "FASE_VAZIA")
        if r["PROJETOS"] == "SP":
            p(r, "PROJETO_PREENCHIDO_COM_UF", "SP")
        if r["PROJETOS"] and re.search(r"S[^a ]o Paulo|SÃo", str(r["PROJETOS"])):
            p(r, "PROJETO_COM_ACENTO_QUEBRADO", r["PROJETOS"])
        if r["FASE"] == "K2552" and r["OPERADORA"] == "VIVO":
            p(r, "OPERADORA_DIVERGE_DA_FASE", "K2552 com VIVO")

    # Duplicidades
    grupos = collections.defaultdict(list)
    for r in linhas:
        grupos[(so_digitos(r["P.O"]), str(r["ITEM"]), chave_projeto(r["PROJETOS"]))].append(r)
    for g in grupos.values():
        if len(g) > 1 and sum(emitida(r) for r in g) >= 2:
            for r in g:
                p(r, "POSSIVEL_FATURAMENTO_EM_DUPLICIDADE",
                  "mesma P.O+item+projeto com outras linhas: " + ",".join(str(x["_linha"]) for x in g if x is not r))
    po_item = collections.Counter((so_digitos(r["P.O"]), str(r["ITEM"])) for r in linhas)
    print(f"\n== Duplicidades")
    print(f"  linhas 100% identicas: {n - len(set(tuple(r[c] for c in colunas) for r in linhas))}")
    print(f"  numeros de P.O repetidos (apos remover apostrofos): "
          f"{sum(1 for v in collections.Counter(so_digitos(r['P.O']) for r in linhas).values() if v > 1)}")
    print(f"  P.O+ITEM repetidos: {sum(1 for v in po_item.values() if v > 1)} grupos")
    dup = [g for g in grupos.values() if len(g) > 1 and sum(emitida(r) for r in g) >= 2]
    print(f"  P.O+ITEM+PROJETO com 2+ notas emitidas: {len(dup)} grupos, "
          f"R$ {sum((r['PREÇO ORIGINAL'] or 0) for g in dup for r in g if emitida(r)):,.2f}")

    with open(os.path.join(saida, "inconsistencias.csv"), "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=["linha", "tipo", "po", "item", "status", "nfse", "detalhe"], delimiter=";")
        w.writeheader()
        w.writerows(sorted(problemas, key=lambda x: (x["tipo"], x["linha"])))
    print("\n== Inconsistencias por tipo (detalhe em inconsistencias.csv)")
    for t, q in collections.Counter(x["tipo"] for x in problemas).most_common():
        print(f"  {q:6}  {t}")


if __name__ == "__main__":
    main()
